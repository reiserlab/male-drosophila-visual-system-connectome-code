# ---
# jupyter:
#   jupytext:
#     text_representation:
#       extension: .py
#       format_name: percent
#       format_version: '1.3'
#       jupytext_version: 1.16.7
#   kernelspec:
#     display_name: default
#     language: python
#     name: python3
# ---

# %%
from pathlib import Path
import re
from datetime import datetime

import pandas as pd
from dotenv import find_dotenv
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

from neuprint import NeuronCriteria as NC, fetch_neurons

from utils import olc_client

from utils.ol_color import OL_COLOR
from utils.ol_types import OLTypes

from queries.webpage_queries import consensus_nt_for_instance
from html_pages.patterns import shorten_nt_name

PROJECT_ROOT = Path(find_dotenv()).parent
c = olc_client.connect(verbose=True)
print(f"Project root directory: {PROJECT_ROOT}")


# %%
# Get the full table
ol = OLTypes()
types = ol.get_neuron_list(side='both')
types["instance"] = types[["type","hemisphere"]].agg("_".join, axis=1)
cell_instances = types['instance']

# Get counts
neurons_df,roi_counts_df = fetch_neurons(NC(instance=cell_instances))
ncells_df = neurons_df.groupby('instance')['bodyId'].nunique().reset_index(name='n_cells')
ncells_sorted_df = ncells_df.sort_values(by='n_cells',ascending=False)
ncells_sorted_df.columns = ['instance','n_cells']


# %%
# Get neuro
# transmitter predictions

# Define a wrapper function for consensus_nt_for_instance to handle single values
def get_nt_prediction(instance):
    nt_prediction = consensus_nt_for_instance(instance=instance)
    # Directly return 'unclear' if that's the result from consensus_nt_for_instance
    if nt_prediction == 'unclear':
        return 'unclear'
    elif nt_prediction is None:
        return "No NT Prediction"
    else:
        return shorten_nt_name(nt_prediction)

# Apply the wrapper function to the 'instance' column to create a new 'nt_prediction' column
types['nt_prediction'] = types['instance'].apply(get_nt_prediction)


# %%
# Define the custom sorting key function
def create_sort_key(s):
    # Extract leading non-numeric part and numeric part
    match = re.match(r"([a-zA-Z]*)(\d*)", s)
    prefix = match.group(1)  # Non-numeric prefix
    number = match.group(2)  # Numeric part

    # Convert numeric part to integer, default to 0 if empty
    number = int(number) if number else 0

    return (prefix.lower(), number)

# Combine and sort the dataframe

# Make table from selected columns
to_table_1 = types[['type','instance', 'main_groups', 'star_neuron', 'nt_prediction']].copy()

# Add n_cells
to_table_1_with_ncells = pd.merge(to_table_1, ncells_sorted_df, on='instance', how='left')

# Create a temporary sorting key for the '_L' and '_R' suffix
to_table_1_with_ncells['sort_key_suffix'] = to_table_1_with_ncells['instance'].str.endswith('_L')

# Apply the sorting key to 'type'
to_table_1_with_ncells['sort_key_type'] = to_table_1_with_ncells['type'].apply(create_sort_key)

# Sort by the temporary keys
table_1_sorted = to_table_1_with_ncells.sort_values(by=['sort_key_suffix', 'sort_key_type', 'instance'], ascending=[True, True, True])

# Drop the temporary sorting key columns
table_1_sorted = table_1_sorted.drop(columns=['sort_key_suffix', 'sort_key_type'])

# Replace the values in the 'main_groups' column
replacements = {
    'OL_intrinsic': 'ONIN',
    'OL_connecting': 'ONCN'
}

# Apply the replacements
table_1_sorted['main_groups'] = table_1_sorted['main_groups'].replace(replacements)

# Reorder columns with 'n_cells' as the third column
desired_order = ['type', 'instance', 'n_cells', 'main_groups', 'star_neuron', 'nt_prediction']
table_1_sorted = table_1_sorted.reindex(columns=desired_order)

# # For debugging
# pd.options.display.max_rows
# print(table_1_sorted)

# %%
# Saving the DataFrame to the Excel file
output_path = Path(PROJECT_ROOT, 'results', 'supp_tables')
output_path.mkdir(parents=True, exist_ok=True)

# Get the current date in yyyy-mm-dd format
current_date = datetime.now().strftime('%Y-%m-%d')

# Construct the filename with the date
filename = f"SuppTable1_Cell-types_and_counts_{current_date}.xlsx"

# Define the full path to save the file
excel_file_path = output_path / filename

# Save the DataFrame to the Excel file
table_1_sorted.to_excel(excel_file_path, sheet_name="Sheet1", index=False)

# %%
# Define font colors
color_mapping_groups = {
    'ONIN': OL_COLOR.OL_TYPES.hex[0]   # OL_intrinsic -> ONIN
  , 'ONCN': OL_COLOR.OL_TYPES.hex[1]   # OL_connecting -> ONCN
  , 'VPN': OL_COLOR.OL_TYPES.hex[2]
  , 'VCN': OL_COLOR.OL_TYPES.hex[3]
  , 'other': OL_COLOR.OL_TYPES.hex[4]  # OL/CB other -> other
}

# Convert hex colors to ARGB format (prepend 'FF' for full opacity)
color_mapping_groups = {key: 'FF' + value[1:] for key, value in color_mapping_groups.items()}

# %%
# Format the excel table

# Load the workbook and select the active worksheet
workbook = openpyxl.load_workbook(excel_file_path)
sheet = workbook.active

# Define a border with no sides (i.e., no border)
no_border = Border(
    left=Side(border_style=None),
    right=Side(border_style=None),
    top=Side(border_style=None),
    bottom=Side(border_style=None)
)

# Define the border for the double-line
bottom_double = Side(style='double')
double_border = Border(
    left=Side(border_style=None),
    right=Side(border_style=None),
    top=Side(border_style=None),
    bottom=bottom_double
)

# Apply no borders to all cells
for row in sheet.iter_rows():
    for cell in row:
        cell.border = no_border

# Apply the double-line border to the bottom of the title row (row 1)
for cell in sheet[1]:  # Title row
    cell.border = double_border

# Rename columns
headers = ['cell type', 'instance', 'no. of cells', 'main groups', 'bodyId in figures', 'predicted neurotransmitter']
for col_idx, header in enumerate(headers, start=1):
    sheet.cell(row=1, column=col_idx, value=header)


# Adjust column widths to fit the widest entry
for col in sheet.columns:
    max_length = 0
    column = col[0].column_letter  # Get the column name
    for cell in col:
        if len(str(cell.value)) > max_length:
            max_length = len(cell.value)
    adjusted_width = (max_length + 2)  # Add a little extra space
    sheet.column_dimensions[column].width = adjusted_width


# Apply gray background to the title row
gray_fill = PatternFill(start_color='D3D3D3', end_color='D3D3D3', fill_type='solid')
for cell in sheet[1]:  # The title row
    cell.fill = gray_fill
    cell.font = Font(bold=True)  # Optional: make the header bold

# Define the column index for 'main_groups' (1-based index)
main_groups_col_idx = 4  # Adjust this if needed after column reordering

# Apply text coloring based on the 'main_groups' column
for row in sheet.iter_rows(min_row=2, max_col=sheet.max_column):
    main_group_value = row[main_groups_col_idx - 1].value  # Adjust index for zero-based
    color_hex = color_mapping_groups.get(main_group_value, 'FF000000')  # Default to black if not found
    font_color = Font(color=color_hex)

    # Apply the font color to the entire row
    for cell in row:
        cell.font = font_color

# Define center alignment
center_alignment = Alignment(horizontal='center')

# Center the text in the "No. of cells" and "BodyID in figures" columns
n_cells_col_idx = headers.index('no. of cells') + 1
bodied_in_figures_col_idx = headers.index('bodyId in figures') + 1

for row in sheet.iter_rows(min_row=2, max_col=sheet.max_column):
    row[n_cells_col_idx - 1].alignment = center_alignment
    row[bodied_in_figures_col_idx - 1].alignment = center_alignment

# Save the workbook with the applied changes
workbook.save(excel_file_path)

# %%
